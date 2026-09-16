"""
Renders a tailored-resume dict (as produced by tailor.tailor_resume) into a
.docx that mirrors the styling/layout of a user-supplied template resume.
The template file is opened read-only (never overwritten); its section
settings, headers/footers, and styles are reused, and the body content is
replaced with the tailored output.
"""
import os
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _safe_filename(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", s).strip("_")


def _clear_body(doc: Document):
    """Remove the template body but keep its sectPr (margins, page size,
    headers/footers). Content gets rebuilt in render() afterwards."""
    body = doc.element.body
    sect_prs = [c for c in body if c.tag.endswith("}sectPr")]
    for sp in sect_prs:
        body.remove(sp)
    for child in list(body):
        body.remove(child)
    for sp in sect_prs:
        body.append(sp)


def _add_para(doc, text="", style=None, bold=False, align=None, size=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size is not None:
            r.font.size = size
    if align is not None:
        p.alignment = align
    return p


def _heading(doc, text: str):
    return _add_para(doc, text, style="Body Text", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def render(contact: dict, tailored: dict, out_dir: str, job_title: str, job_company: str, template_path: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    doc = Document(template_path)
    _clear_body(doc)

    # Name (Normal, bold, centered -- as in the template)
    _add_para(doc, contact.get("name", ""), style="Normal", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Contact line (Body Text, centered, " • " separators -- as in the template)
    contact_bits = [contact.get("email", ""), contact.get("phone", ""), contact.get("location", "")]
    for link in contact.get("links", []):
        contact_bits.append(f"{link.get('label', '')}: {link.get('url', '')}")
    contact_line = " • ".join(b for b in contact_bits if b)
    if contact_line:
        _add_para(doc, contact_line, style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph(style="Body Text")  # spacer

    # Summary
    if tailored.get("summary"):
        _heading(doc, "Summary")
        _add_para(doc, tailored["summary"], style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER)

    # Education
    if tailored.get("education"):
        _heading(doc, "Education")
        for edu in tailored["education"]:
            school = edu.get("school", "")
            degree = edu.get("degree", "")
            graduation = edu.get("graduation", "")
            # School name line (Normal, bold)
            school_p = _add_para(doc, school, style="Normal", bold=True)
            # Degree + graduation (Body Text)
            deg_p = doc.add_paragraph(style="Body Text")
            deg_text = degree
            if graduation:
                deg_text += f"\t{graduation}"
            deg_p.add_run(deg_text)

    # Experience
    if tailored.get("experience"):
        _heading(doc, "Experience")
        for job in tailored.get("experience", []):
            company_line = _add_para(doc, job.get("company", ""), style="Normal", bold=True)
            location = job.get("location", "")
            if location:
                company_line.add_run(f"\t{location}")
            dates = " – ".join(filter(None, [job.get("start_date", ""), job.get("end_date", "")]))
            title_line = _add_para(doc, job.get("title", ""), style="Normal", bold=True)
            if dates:
                title_line.add_run(f"\t{dates}")
            for bullet in job.get("bullets", []):
                _add_para(doc, bullet, style="List Paragraph")

    # Skills & Interests -- one section with sub-labels, as in the template.
    # Technical skill groups have a bold label with each item bulleted.
    has_any = (tailored.get("skills") or tailored.get("soft_skills")
               or tailored.get("languages") or tailored.get("interests"))
    if has_any:
        _heading(doc, "Skills & Interests")

    if tailored.get("skills"):
        _add_para(doc, "Technical Skills:", style="Body Text", bold=True)
        for group in tailored.get("skills", []):
            label = group.get("group", "")
            if label:
                p = doc.add_paragraph(style="Body Text")
                p.add_run(f"{label}:").bold = True
            for item in group.get("items", []):
                _add_para(doc, item, style="List Paragraph")

    # Soft Skills
    if tailored.get("soft_skills"):
        _add_para(doc, "Soft Skills:", style="Body Text", bold=True)
        for skill in tailored["soft_skills"]:
            _add_para(doc, skill, style="List Paragraph")

    # Human languages (as in the template's "Language:" block)
    if tailored.get("languages"):
        _add_para(doc, "Language:", style="Body Text", bold=True)
        for lang in tailored["languages"]:
            _add_para(doc, lang, style="Body Text")

    # Interests
    if tailored.get("interests"):
        _add_para(doc, "Interests:", style="Body Text", bold=True)
        for interest in tailored["interests"]:
            _add_para(doc, interest, style="Body Text")

    fname = _safe_filename(f"resume_{job_company}_{job_title}") + ".docx"
    out_path = os.path.join(out_dir, fname)
    doc.save(out_path)
    return out_path
