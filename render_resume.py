"""
Render a tailored-resume dict (as produced by tailor.tailor_resume) into a
clean, single-column .docx -- name header, contact line, then Summary /
Skills / Experience / Education sections.

No template file needed: styling is self-contained (11pt, Calibri-ish
default, centered header, bold section headings), so the dashboard
download endpoint can render from JSON alone.
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _safe_filename(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", s).strip("_")[:80]


def _heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    return p


def _body(doc, text: str, bullet: bool = False, italic: bool = False):
    p = doc.add_paragraph(style="List Bullet" if bullet else "Normal")
    if text:
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.italic = italic
    return p


def render(contact: dict, tailored: dict, out_dir: str, job_title: str,
           job_company: str) -> str:
    """Write the .docx and return its path."""
    os.makedirs(out_dir, exist_ok=True)
    doc = Document()
    for p in doc.paragraphs:  # default template starts with one empty para
        if not p.text.strip():
            p._element.getparent().remove(p._element)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(contact.get("name", ""))
    run.bold = True
    run.font.size = Pt(14)

    bits = [contact.get("email", ""), contact.get("phone", ""),
            contact.get("location", "")]
    bits += [f"{l.get('label', '')}: {l.get('url', '')}"
             for l in contact.get("links", [])]
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run(" • ".join(b for b in bits if b))
    run.font.size = Pt(10)

    if tailored.get("summary"):
        _heading(doc, "Summary")
        _body(doc, tailored["summary"])

    if tailored.get("skills"):
        _heading(doc, "Skills")
        for g in tailored["skills"]:
            items = g.get("items", [])
            if not items:
                continue
            p = doc.add_paragraph()
            run = p.add_run(f"{g.get('group', 'Skills')}: ")
            run.bold = True
            run.font.size = Pt(11)
            run = p.add_run(", ".join(items))
            run.font.size = Pt(11)

    if tailored.get("experience"):
        _heading(doc, "Experience")
        for job in tailored["experience"]:
            p = doc.add_paragraph()
            run = p.add_run(
                f"{job.get('title', '')} — {job.get('company', '')}")
            run.bold = True
            run.font.size = Pt(11)
            meta = " | ".join(b for b in [
                job.get("location", ""),
                " – ".join(d for d in [job.get("start_date", ""),
                                       job.get("end_date", "")] if d),
            ] if b)
            if meta:
                _body(doc, meta, italic=True)
            for b in job.get("bullets", []):
                _body(doc, b, bullet=True)

    bank_edu = tailored.get("education")
    if bank_edu:
        _heading(doc, "Education")
        for e in bank_edu:
            if isinstance(e, dict):
                _body(doc, f"{e.get('degree', '')}, {e.get('school', '')} "
                           f"({e.get('graduation', '')})".strip(" ,()"))
            else:
                _body(doc, str(e))

    if tailored.get("certifications"):
        _heading(doc, "Certifications")
        for c in tailored["certifications"]:
            _body(doc, c if isinstance(c, str) else str(c), bullet=True)

    for key, title in (("soft_skills", "Soft Skills"),
                       ("languages", "Languages"),
                       ("interests", "Interests")):
        if tailored.get(key):
            _heading(doc, title)
            _body(doc, ", ".join(tailored[key]))

    fname = _safe_filename(
        f"Resume_{job_title}_at_{job_company}.docx") or "Resume_tailored.docx"
    path = os.path.join(out_dir, fname)
    doc.save(path)
    return path
